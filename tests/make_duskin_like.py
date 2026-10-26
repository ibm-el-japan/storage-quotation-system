"""新様式(A1:F85)の見積もり依頼記入シートを、架空の値で再現した検証用xlsxを作る (t15 の入力)。
実案件のシート(顧客情報を含む)は公開できないため、レイアウト・ラベル・チェック記号・セル番地だけを同じにして、値はダミーにしている。
行番号は旧テンプレート(A1:F91)から1〜6行ずれる: 記入日 C4 / 案件名 C10 / Opp(HW/SW) C12 / ストレージ担当営業 C15 / Expert Labs担当営業 C17 /
Offering C20-C25 / 支援対象製品 C27-C40 / 支援期間 C41 / 支援(作業実施)場所 C42 / 支援内容(詳細) C43 / 事前確認 D51 / 提供形態 C70-C73 /
契約形態 C75-C78 / 支払い計画 C80-C82 / 月次作業報告書 C83。"""
import os, datetime
import openpyxl

here = os.path.dirname(os.path.abspath(__file__))
dst = os.path.join(here, 'duskin_like.xlsx')
wb = openpyxl.Workbook()
ws = wb.active
ws.title = '見積もり依頼記入シート'
S = {}
def put(ref, v): S[ref] = v; ws[ref] = v

put('A1', 'Storage Expert Labs支援依頼フォーム')
put('B4', '記入日*');                         ws['C4'] = datetime.datetime(2026, 10, 26); ws['C4'].number_format = 'mm-dd-yy'
put('B5', '記入者*');                         put('C5', '記入 太郎(Kinyu Taro)')
put('A7', '案件情報')
put('B8', '確認事項'); put('C8', '記入欄'); put('D8', '備考')
put('B9', 'お客様名*');                       put('C9', 'サンプル商事様')
put('B10', '案件名*');                        put('C10', 'バックアップ基盤構築')
put('B11', 'DNSO #')
put('B12', 'Opportunity No（HW/SW製品）*');   put('C12', '006xx0000000001AAA\n ')
put('B13', 'Opportunity No（サービス）*');    put('D13', 'Opportunity No（HW/SW製品）への登録でも可')
put('B14', 'ビジネス・パートナー様名');       put('C14', 'ダミーBP株式会社'); put('D14', 'BP商流の場合')
put('B15', 'ストレージ担当営業名*');          put('C15', '営業 一郎(Eigyo Ichiro)'); put('D15', '未定の場合TBD')
put('B16', 'ストレージ担当テクニカル・セールス名*'); put('C16', '技術 二郎(Gijutsu Jiro)'); put('D16', '未定の場合TBD')
put('B17', 'Expert Labs担当営業名*');         put('D17', '未定の場合TBD')
put('B18', 'その他案件関係者名');             put('D18', 'SSR、AP、Technology Sales等')
put('B19', 'Offering*');                      put('C19', '支援対象Offeringを選択（左端の+（26行目横）をクリック）')
for r, t in zip(range(20, 26), ['□Power Expert Labs',
                                '☑️Storage Expert Labs（Product Code：30C9J/Product Name：Expert Labs for Storage）',
                                '□zHW Expert Labs', '□zTPS Expert Labs', '□zHW GDPS', '□その他（    )']):
    put('C%d' % r, t)
put('B26', '支援対象製品*');                  put('C26', '支援対象製品を選択（左端の+（41行目横）をクリック）')
prods = ['□Storage Virtualize系製品（□SVC□FlashSystem xxxx□SV4PC）', '□DS8K（製品名：DS8xxx）', '□Storage Scale System（製品名：SSSxxxx）',
         '□テープ製品（製品名：TSxxxx）', '□SAN Switch（□b-type（SANxxB-x）□c-type（SANxxC-x））', '□Storage Scale', '□Storage Protect',
         '□Spectrum Control', '□Storage Insights', '□Storage Ceph（□Ready Nodes□Software）', '□IBM Fusion（□HCI□Software）',
         '☑️Storage Defender（☑️Data Protect□Sentinel□DRS/DMS）', '□Cloud Object Storage（□Appliance□Software）', '□その他（xxxxx（対象製品名を追記））']
for i, t in enumerate(prods): put('C%d' % (27 + i), t)
put('B41', '支援期間*');                      put('C41', '2027年9月～2027年11月30日')
put('B42', '支援(作業実施)場所*');            put('C42', '大阪')
put('B43', '支援内容（詳細）*')
put('C43', '提出ドキュメント\n\u3000・基本設計書\t\n\u3000・詳細設計書（パラメーターシート）\t\n\u3000・テスト計画書兼報告書\nHW製品構築、SW製品導入構築\nQA支援')
put('D43', '例） 下記項目を中心に、台数や規模などの情報も記載')
put('B44', '機密情報取扱い*');                put('C44', '☑️無し\n□有り'); put('D44', '機密情報取扱い有りの場合は、機密保持契約締結情報等を記入')
put('B45', '社内関係者との打ち合わせの調整*'); put('C45', '☑️調整します')
put('B46', '社外関係者との打ち合わせの調整*'); put('C46', '☑️調整します（ただ、現時点で詳細が未定のため、概算だけお願いします）')
put('A48', '事前確認事項（要記入）')
put('B49', '事前確認事項（左端の+（64行目横）をクリック）'); put('D49', '記入欄（Y / N）'); put('E49', '備考')
put('B50', '支援対象製品/構成/機能について、提案時に製品有識者への事前相談は提案内容の精度を高める意味があります。')
put('B51', '\u3000SO案件の場合、「SOセールスにおけるSTGテクニカル支援依頼プロセス」を活用されていますか？\u3000→'); put('D51', 'N'); put('E51', 'SO案件ではない')
put('B52', '\u3000上記プロセスを活用されている場合、相談者名を記入ください。↓')
put('B53', '相談者名：')
put('B54', '支援対象製品/構成/機能の妥当性確認について、提案時における適切なプロセスの実施がサービス品質の向上に繋がります。')
for r, t in zip(range(55, 63), ['提案内容の妥当性を確認するために、SAレビューやTCTを実施されていますか？',
                                 'SAレビューやTCT以外の提案時に必要なレビューを実施されていますか？',
                                 '提案内容の妥当性を確認する上で、各製品で提供されている提案前チェックリストを活用されていますか？',
                                 'パフォーマンスの妥当性について、事前にシミュレーションツールなどは活用されていますか？',
                                 'キャパシティーの妥当性について、事前にシミュレーションツールなどは活用されていますか？',
                                 '支援対象構成について、適切なサポート条件の情報を確認されていますか？',
                                 '支援対象製品について、制限/前提事項の情報を確認されていますか？',
                                 '正式サポートされていない構成や使用を想定されている場合、事前に例外承認は取得されていますか？']):
    put('B%d' % r, '\u3000' + t)
put('B63', '支援契約終了後のサポートについて、標準保守契約以外の追加契約の事前検討はサポート品質を維持する上での重要事項となります。')
put('B64', '\u3000見積もりいただいている支援契約が終了した後に、適切な保守契約を検討いただいていますでしょうか？')
put('B65', '\u3000見積もりいただいている支援契約が終了した後に、適切な運用支援サービス契約※1を検討いただいていますでしょうか？')
put('A67', '契約情報'); put('C67', '*記入必須')
put('B68', '確認事項'); put('C68', '記入欄'); put('D68', '備考')
put('B69', 'ご提供形態*');                    put('C69', '支援のご提供形態を選択（左端の+（72行目横）をクリック）')
put('C70', '☑️支援契約（Expert Labsの単独契約）')
put('C71', '□支援契約（他サービス部門とSLINKでのワーク分割）')
put('C72', '□他部門のSLINK案件に対するDOUによる作業支援\u3000（レイバーはILC計上）※2')
put('C73', '□その他（xxxxx（提供形態を追記））')
put('B74', '契約形態');                       put('C74', '契約形態を選択（左端の+（77行目横）をクリック）')
put('C75', '□直販                     （例）\u3000エンドユーザ様-IBM')
put('C76', '☑️直販（サブコン）\u3000\u3000（例）\u3000エンドユーザ様-BP様-IBM')
put('C77', '□BP商流'); put('C78', '□その他 （xxxxx（契約形態を追記））')
put('B79', '支払い計画');                     put('C79', '支払い計画を選択（左端の+（81行目横）をクリック）')
put('C80', '☑️各月（標準）'); put('C81', '□支払い一回（支援期間3ヶ月以内のみ選択可）'); put('C82', '□その他（xxxxx（支払い計画を追記））')
put('B83', '月次作業報告書提出の必要性');     put('C83', '□有 ☑️無')
put('B85', '※ レイバー（工数）はTime@IBMで計上、交通費が発生する場合はTime@IBM向けのプロジェクトコードで実費精算。')
wb.save(dst)
print('wrote', dst)
